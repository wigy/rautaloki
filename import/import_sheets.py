# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl"]
# ///
"""One-off import of the old Google Sheets workbooks into Rautaloki JSON files.

    uv run import/import_sheets.py            # writes data/
    uv run import/import_sheets.py --check    # only report, write nothing

Sources are the xlsx exports in import/xlsx/: master.xlsx (Weight Master: catalog,
programs, weight tables) and one workbook per user, <user id>.xlsx (Scores sheet =
training cycles, name in A1; Records sheet = cross-check only, since records are
derived from cycles).
"""

import datetime as dt
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "import" / "xlsx"
DATA = ROOT / "data"
PLATES = [1, 1.25, 2.25, 2.5, 5, 10, 15]

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
problems: list[str] = []


def slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    try:
        return num(float(str(v).replace(",", ".").replace("kg", "").strip()))
    except ValueError:
        return None


def date(v, where: str):
    """Sheets parsed d/m/y input as m/d/y whenever the day was <= 12, so every real
    date value has its day and month swapped. Text values are still d/m/y as typed."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return dt.date(v.year, v.day, v.month).isoformat()
    m = re.fullmatch(r"(\d{1,2})[/.]+(\d{1,2})[/.]+(\d{2,4})", str(v).strip())
    if not m:
        problems.append(f"{where}: unparseable date {v!r}, dropped")
        return None
    d, mo, y = (int(x) for x in m.groups())
    if y < 100:
        y += 2000
    return dt.date(y, mo, d).isoformat()


def rows(ws, start: int):
    for i, row in enumerate(ws.iter_rows(min_row=start, values_only=True), start):
        yield i, row


def load_catalog(wb):
    exercises = []
    for i, r in rows(wb["Catalog"], 3):
        muscle, name, equipment, level, other, url = (r + (None,) * 6)[:6]
        if not name:
            continue
        split = lambda s: [x.strip() for x in str(s).split(",") if x.strip()] if s else []
        exercises.append({
            "id": slug(name),
            "name": name.strip(),
            "mainMuscle": muscle.strip(),
            "otherMuscles": split(other),
            "equipment": split(equipment),
            "level": level,
            "url": url,
        })
    return exercises


def load_programs(wb):
    programs = []
    for i, r in rows(wb["Programs"], 3):
        if not r[0]:
            continue
        weeks = []
        for w in range(10):
            scheme, pct = r[1 + w], r[11 + w]
            if not scheme:
                break
            weeks.append({"scheme": str(scheme).strip(), "pct": round(float(pct), 4)})
        programs.append({"id": slug(r[0]), "name": r[0], "weeks": weeks})
    return programs


def load_weights(wb):
    """Achievable loads per equipment kind, as plate counts per side/end. The sheet's
    'Multiple' says how many of each listed count are actually loaded: 4 for a pair of
    dumbbells, 2 for one dumbbell or a barbell, 1 for other (cable, T-bar)."""
    kinds = {
        "dumbbells": ("Dumbbells", 1, 4),
        "dumbbell": ("Dumbbells", 1, 2),
        "barbell": ("Barbell", 7, 2),
        "other": ("Other", 0, 1),
    }
    out = {}
    for kind, (sheet, base, multiple) in kinds.items():
        loads = []
        for i, r in rows(wb[sheet], 3):
            total = num(r[8])
            if total is None:
                continue
            plates = {str(p): int(r[k]) for k, p in enumerate(PLATES) if num(r[k])}
            if not plates and base:
                total = base  # the "- (7kg)" row: empty bar / bare handle
            loads.append({"total": total, "plates": plates})
        out[kind] = {"base": base, "multiple": multiple, "loads": loads}
    return out


def load_cycles(wb, user: str, programs: dict, catalog: dict):
    ws = wb["Scores"]
    cycles = []
    for i, r in rows(ws, 4):
        r = (r + (None,) * 20)[:20]
        name = r[0]
        if not name:
            continue
        where = f"{user} Scores!A{i} {name}"
        prog = programs.get(slug(r[1] or ""))
        if not prog:
            problems.append(f"{where}: unknown program {r[1]!r}, row skipped")
            continue
        if slug(name) not in catalog:
            problems.append(f"{where}: exercise not in catalog")
        n = len(prog["weeks"])
        scores = []
        for w, v in enumerate(r[7:17]):
            s = "-" if isinstance(v, str) and v.strip() == "-" else num(v)
            if w >= n:
                if s not in (None, "-"):
                    problems.append(f"{where}: score {s} in W{w+1} beyond {n}-week program")
                continue
            scores.append(s)
        add = num(r[3])
        if r[3] not in (None, "") and add is None:
            problems.append(f"{where}: non-numeric Add? {r[3]!r}, ignored")
        started = date(r[5], where)
        finished = date(r[18], where)
        if started and finished and finished < started:
            problems.append(f"{where}: finished {finished} before started {started}")
        cycle = {
            "id": f"{started or 'undated'}-{slug(name)}",
            "exercise": slug(name),
            "program": prog["id"],
            "started": started,
            "oldBest": num(r[6]),
            "scores": scores,
            "add": add or 0,
            "done": str(r[2] or "").strip().upper() == "F" or finished is not None,
            "finished": finished,
            "newBest": num(r[17]),
        }
        ids = {c["id"] for c in cycles}
        k = 2
        while cycle["id"] in ids:
            cycle["id"] = f"{started}-{slug(name)}-{k}"
            k += 1
        cycles.append(cycle)
    return cycles


def check_records(wb, user, cycles):
    best: dict[str, float] = {}
    for c in cycles:
        for v in (c["oldBest"], c["newBest"]):
            if isinstance(v, (int, float)):
                best[c["exercise"]] = max(best.get(c["exercise"], 0), v)
    for i, r in rows(wb["Records"], 4):
        if len(r) > 2 and r[1] and num(r[2]) is not None:
            got = best.get(slug(r[1]))
            if got != num(r[2]):
                problems.append(f"{user} Records!C{i} {r[1]}: sheet {num(r[2])}, derived {got}")


def write(path: Path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=False) + "\n")


def main():
    check = "--check" in sys.argv
    master = openpyxl.load_workbook(XLSX / "master.xlsx", data_only=True)
    catalog = load_catalog(master)
    programs = load_programs(master)
    weights = load_weights(master)
    by_id = {p["id"]: p for p in programs}
    cat_ids = {e["id"]: e for e in catalog}
    users = {}
    for path in sorted(XLSX.glob("*.xlsx")):
        uid = path.stem
        if uid == "master":
            continue
        wb = openpyxl.load_workbook(path, data_only=True)
        uname = str(wb["Scores"]["A1"].value or uid.capitalize()).strip()
        cycles = load_cycles(wb, uid, by_id, cat_ids)
        check_records(wb, uid, cycles)
        users[uid] = {"id": uid, "name": uname, "cycles": cycles}
        active = sum(1 for c in cycles if not c["done"])
        print(f"{uname}: {len(cycles)} cycles, {active} active")
    print(f"catalog {len(catalog)} exercises, {len(programs)} programs")
    for p in problems:
        print("  !", p)
    if check:
        return
    write(DATA / "catalog.json", catalog)
    write(DATA / "programs.json", programs)
    write(DATA / "weights.json", weights)
    for uid, u in users.items():
        write(DATA / "users" / f"{uid}.json", u)
    print(f"wrote {DATA}")


main()
